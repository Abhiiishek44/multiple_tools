import json
import tempfile
from collections.abc import Iterator
from pathlib import Path
from typing import BinaryIO
from uuid import uuid4

from docx import Document

from infrastructure.queue import AI_OCR_QUEUE, CHAT_INGEST_TASK, celery_app
from packages.chat import provider, repository
from packages.chat.models import ChatDocument, Citation, Conversation, Message
from packages.core.config import get_settings
from packages.core.errors import ConflictError, NotFoundError, ValidationError
from packages.documents.pdf import convert_pdf_to_text
from packages.documents.text import TextNormalizer
from packages.storage import get_storage

SUPPORTED_INPUTS = {
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain"},
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
}


def create_conversation(user_id: str, title: str) -> Conversation:
    normalized = title.strip()
    if not normalized:
        raise ValidationError("Conversation title is required")
    return repository.create_conversation(uuid4().hex, user_id, normalized[:200])


def get_conversation(
    conversation_id: str, user_id: str
) -> tuple[Conversation, list[ChatDocument], list[Message]]:
    conversation = repository.get_conversation(conversation_id, user_id)
    if conversation is None:
        raise NotFoundError("Conversation not found")
    return (
        conversation,
        repository.list_documents(conversation_id),
        repository.list_messages(conversation_id),
    )


def submit_document(
    *,
    conversation_id: str,
    user_id: str,
    filename: str,
    media_type: str,
    stream: BinaryIO,
) -> ChatDocument:
    if repository.get_conversation(conversation_id, user_id) is None:
        raise NotFoundError("Conversation not found")
    suffix = Path(filename).suffix.lower()
    normalized_media_type = media_type.partition(";")[0].strip().lower()
    if suffix not in SUPPORTED_INPUTS or normalized_media_type not in SUPPORTED_INPUTS[suffix]:
        raise ValidationError("Chat supports TXT, Markdown, PDF, and DOCX documents")
    document_id = uuid4().hex
    artifact_key = f"chat/{conversation_id}/{document_id}/input{suffix}"
    storage = get_storage()
    storage.save_stream(artifact_key, stream, get_settings().max_upload_bytes)
    try:
        document = repository.create_document(
            document_id,
            conversation_id,
            Path(filename).name,
            normalized_media_type,
            artifact_key,
        )
        celery_app.send_task(
            CHAT_INGEST_TASK,
            args=[document.id],
            task_id=f"chat-document-{document.id}",
            queue=AI_OCR_QUEUE,
            ignore_result=True,
        )
        return document
    except Exception:
        storage.delete(artifact_key)
        raise


def ingest_document(document_id: str) -> None:
    document = repository.get_document(document_id)
    if document is None:
        raise ValueError(f"Chat document not found: {document_id}")
    if not repository.mark_document_processing(document.id):
        return
    try:
        with tempfile.TemporaryDirectory(prefix=f"chat-{document.id}-") as directory:
            source = Path(directory) / f"input{Path(document.filename).suffix.lower()}"
            get_storage().download_file(document.artifact_key, source)
            text = _extract_text(source, document.id)
        chunks = _chunk_text(TextNormalizer().normalize(text))
        if not chunks:
            raise ValueError("The document contains no extractable text")
        repository.replace_chunks_and_mark_ready(document.id, chunks)
    except Exception as error:
        repository.mark_document_failed(document.id, str(error))
        raise


def answer_question(conversation_id: str, user_id: str, question: str) -> Message:
    messages, citations = _prepare_question(conversation_id, user_id, question)
    answer = provider.complete(messages)
    return _save_answer(conversation_id, answer, citations)


def stream_answer(
    conversation_id: str, user_id: str, question: str
) -> Iterator[str]:
    messages, citations = _prepare_question(conversation_id, user_id, question)
    citation_payload = [_citation_payload(citation) for citation in citations]
    yield _sse("citations", citation_payload)
    parts: list[str] = []
    try:
        for token in provider.stream(messages):
            parts.append(token)
            yield _sse("token", {"content": token})
        message = _save_answer(conversation_id, "".join(parts), citations)
        yield _sse(
            "done",
            {"message_id": message.id, "created_at": message.created_at.isoformat()},
        )
    except Exception as error:
        yield _sse("error", {"detail": str(error)})


def _prepare_question(
    conversation_id: str, user_id: str, question: str
) -> tuple[list[dict[str, str]], list[Citation]]:
    conversation, documents, history = get_conversation(conversation_id, user_id)
    del conversation
    normalized = question.strip()
    if not normalized:
        raise ValidationError("Message content is required")
    if not any(document.status == "READY" for document in documents):
        raise ConflictError("At least one document must finish processing before chatting")
    citations = repository.retrieve_chunks(conversation_id, normalized)
    context = "\n\n".join(
        f"[Source {index}: {citation.filename}, chunk {citation.chunk_index}]\n{citation.excerpt}"
        for index, citation in enumerate(citations, start=1)
    )
    repository.create_message(uuid4().hex, conversation_id, "user", normalized)
    messages = [
        {"role": message.role, "content": message.content}
        for message in history[-10:]
    ]
    messages.append(
        {
            "role": "user",
            "content": f"Document excerpts:\n{context}\n\nQuestion: {normalized}",
        }
    )
    return messages, citations


def _save_answer(
    conversation_id: str, answer: str, citations: list[Citation]
) -> Message:
    citation_payload = [_citation_payload(citation) for citation in citations]
    return repository.create_message(
        uuid4().hex,
        conversation_id,
        "assistant",
        answer,
        citation_payload,
    )


def _extract_text(source: Path, document_id: str) -> str:
    suffix = source.suffix.lower()
    if suffix in {".txt", ".md"}:
        return source.read_text(encoding="utf-8")
    if suffix == ".docx":
        document = Document(source)
        blocks = [paragraph.text for paragraph in document.paragraphs]
        blocks.extend(
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        return "\n".join(blocks)
    if suffix == ".pdf":
        destination = source.with_suffix(".txt")
        convert_pdf_to_text(
            source,
            destination,
            job_id=f"chat-document-{document_id}",
            report_progress=lambda _: None,
        )
        return destination.read_text(encoding="utf-8")
    raise ValueError(f"Unsupported chat document type: {suffix}")


def _chunk_text(text: str, size: int = 3000, overlap: int = 300) -> list[str]:
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            boundary = text.rfind("\n", start + size // 2, end)
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def _citation_payload(citation: Citation) -> dict[str, object]:
    return {
        "document_id": citation.document_id,
        "filename": citation.filename,
        "chunk_index": citation.chunk_index,
        "excerpt": citation.excerpt[:500],
    }


def _sse(event: str, payload: object) -> str:
    return f"event: {event}\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"
