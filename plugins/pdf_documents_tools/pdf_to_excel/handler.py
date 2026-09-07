from pathlib import Path

from plugins.base import ToolContext


def convert(context: ToolContext, source: Path, destination: Path) -> Path:
    import pdfplumber
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    with pdfplumber.open(source) as document:
        if not document.pages:
            raise ValueError("PDF contains no pages")
        for page_number, page in enumerate(document.pages, start=1):
            sheet = workbook.create_sheet(f"Page {page_number}")
            tables = page.extract_tables()
            if tables:
                row_number = 1
                for table in tables:
                    for row in table:
                        for column, value in enumerate(row, start=1):
                            sheet.cell(row_number, column, value)
                        row_number += 1
                    row_number += 1
            else:
                for row_number, line in enumerate((page.extract_text() or "").splitlines(), start=1):
                    sheet.cell(row_number, 1, line)
            context.report_progress(10 + int(80 * page_number / len(document.pages)))
    workbook.save(destination)
    return destination
