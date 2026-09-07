import csv
from pathlib import Path

from openpyxl import load_workbook

from plugins.base import ToolContext


def convert(context: ToolContext, source: Path, destination: Path) -> Path:
    context.report_progress(10)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("Workbook contains no worksheets")
        with destination.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])
    finally:
        workbook.close()
    context.report_progress(90)
    return destination
