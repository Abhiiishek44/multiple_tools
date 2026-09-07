from pathlib import Path

from docx import Document

from plugins.base import ToolContext


def convert(context: ToolContext, source: Path, destination: Path) -> Path:
    context.report_progress(20)
    document = Document(source)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    destination.write_text("\n".join(blocks) + "\n", encoding="utf-8")
    context.report_progress(90)
    return destination
