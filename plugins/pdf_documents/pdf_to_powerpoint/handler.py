import tempfile
from pathlib import Path

from plugins.base import ToolContext


def convert(context: ToolContext, source: Path, destination: Path) -> Path:
    import pymupdf
    from pptx import Presentation

    presentation = Presentation()
    with pymupdf.open(source) as document:
        if document.page_count == 0:
            raise ValueError("PDF contains no pages")
        first_page = document[0].rect
        presentation.slide_width = int(first_page.width * 12700)
        presentation.slide_height = int(first_page.height * 12700)
        with tempfile.TemporaryDirectory(prefix="pdf-pptx-") as directory:
            for index, page in enumerate(document):
                image = Path(directory) / f"page-{index + 1}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).save(image)
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                slide.shapes.add_picture(
                    str(image), 0, 0,
                    width=presentation.slide_width,
                    height=presentation.slide_height,
                )
                context.report_progress(10 + int(80 * (index + 1) / document.page_count))
    presentation.save(destination)
    return destination
